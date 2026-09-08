import os
import re
import uuid
import json
import csv
from typing import List, Dict, Any, Tuple, Optional
from psycopg import connect
from dotenv import load_dotenv

# Ensure environment is loaded
workspace_env_path = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), ".env")
load_dotenv(dotenv_path=workspace_env_path)

def get_db_connection():
    """Establishes connection to the PostgreSQL database."""
    return connect(
        host=os.getenv("DB_HOST", "localhost"),
        port=os.getenv("DB_PORT", "5432"),
        user=os.getenv("DB_USER", "postgres"),
        password=os.getenv("DB_PASSWORD"),
        dbname=os.getenv("DB_NAME", "querypilot")
    )

def ensure_metadata_table():
    """Creates the datasets_metadata table if it does not exist."""
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""
            CREATE TABLE IF NOT EXISTS datasets_metadata (
                dataset_id VARCHAR(50) PRIMARY KEY,
                session_id VARCHAR(100) NOT NULL,
                original_filename VARCHAR(255) NOT NULL,
                table_name VARCHAR(100) NOT NULL,
                columns_metadata JSONB NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            """)
            conn.commit()
    finally:
        conn.close()

def sanitize_identifier(name: str) -> str:
    """Sanitizes file and column names to prevent SQL injections or syntax issues."""
    # Convert to lowercase, replace non-alphanumeric chars with underscores
    clean = re.sub(r'[^a-zA-Z0-9_]', '_', name.strip().lower())
    # Ensure it starts with a letter
    if not clean or not clean[0].isalpha():
        clean = "col_" + clean
    # Prevent double underscores and trailing underscores
    clean = re.sub(r'_+', '_', clean).strip('_')
    return clean

def infer_postgres_type(values: List[str]) -> str:
    """Infers PostgreSQL data type based on a list of string values."""
    if not values:
        return "VARCHAR(255)"
        
    # Check for integer
    is_int = True
    for val in values:
        if not val:
            continue
        try:
            int(val.replace(',', ''))
        except ValueError:
            is_int = False
            break
            
    if is_int:
        return "INTEGER"
        
    # Check for numeric/float
    is_float = True
    for val in values:
        if not val:
            continue
        try:
            float(val.replace(',', ''))
        except ValueError:
            is_float = False
            break
            
    if is_float:
        return "DECIMAL(18, 4)"
        
    # Check for date (very simple check for YYYY-MM-DD or standard formats)
    is_date = True
    date_patterns = [
        r'^\d{4}-\d{2}-\d{2}$',
        r'^\d{2}/\d{2}/\d{4}$',
        r'^\d{4}/\d{2}/\d{2}$'
    ]
    for val in values:
        if not val:
            continue
        matched = False
        for pattern in date_patterns:
            if re.match(pattern, val.strip()):
                matched = True
                break
        if not matched:
            is_date = False
            break
            
    if is_date:
        return "DATE"
        
    return "VARCHAR(255)"

def get_dataset_table_name(dataset_id: str) -> str:
    """Retrieves the physical SQL table name for a dataset ID."""
    ensure_metadata_table()
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT table_name FROM datasets_metadata WHERE dataset_id = %s;", (dataset_id,))
            row = cur.fetchone()
            if not row:
                raise ValueError(f"Dataset {dataset_id} does not exist.")
            return row[0]
    finally:
        conn.close()

def get_dataset_ddl(dataset_id: str) -> str:
    """Dynamically generates the DDL description for prompting the Text-to-SQL agent."""
    ensure_metadata_table()
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""
                SELECT table_name, original_filename, columns_metadata 
                FROM datasets_metadata WHERE dataset_id = %s;
            """, (dataset_id,))
            row = cur.fetchone()
            if not row:
                raise ValueError(f"Dataset {dataset_id} does not exist.")
                
            table_name, filename, cols_meta_json = row
            cols_meta = json.loads(cols_meta_json) if isinstance(cols_meta_json, str) else cols_meta_json
            
            ddl_lines = [f"-- Table: {table_name} (Source: {filename})"]
            ddl_lines.append(f"CREATE TABLE {table_name} (")
            
            for orig, sanitized, db_type in cols_meta:
                ddl_lines.append(f"    {sanitized} {db_type}, -- Original Column Name: \"{orig}\"")
                
            # Close the mock DDL string
            if len(ddl_lines) > 2:
                ddl_lines[-1] = ddl_lines[-1].rstrip(',')
            ddl_lines.append(");")
            
            return "\n".join(ddl_lines)
    finally:
        conn.close()

def process_and_save_dataset(
    file_bytes: bytes, 
    filename: str, 
    session_id: str
) -> str:
    """
    Parses a CSV or Excel file, infers its schema, creates a dynamic table,
    performs bulk inserts, and records metadata in PostgreSQL.
    """
    ensure_metadata_table()
    
    # 1. Parse File Content
    rows: List[List[str]] = []
    if filename.endswith(".csv"):
        content = file_bytes.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(content.splitlines())
        rows = list(reader)
    elif filename.endswith(".xlsx"):
        # Excel parsing via openpyxl
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            # Convert values to strings
            rows.append([str(val) if val is not None else "" for val in row])
    else:
        raise ValueError("Unsupported file format. Only CSV and XLSX are allowed.")
        
    if not rows or len(rows) < 2:
        raise ValueError("File is empty or lacks data rows.")
        
    # 2. Extract and Sanitize Headers
    original_headers = [h.strip() for h in rows[0]]
    if not any(original_headers):
        raise ValueError("Header row is empty or invalid.")
        
    sanitized_headers: List[str] = []
    seen = set()
    for idx, header in enumerate(original_headers):
        name = header if header else f"unnamed_column_{idx+1}"
        sanitized = sanitize_identifier(name)
        # Avoid duplicate names by appending suffix
        count = 1
        original_sanitized = sanitized
        while sanitized in seen:
            sanitized = f"{original_sanitized}_{count}"
            count += 1
        seen.add(sanitized)
        sanitized_headers.append(sanitized)
        
    # 3. Sample Data for Type Inference
    data_rows = rows[1:]
    num_cols = len(original_headers)
    
    # Clean rows to match header count
    cleaned_data_rows = []
    for r in data_rows:
        if not any(r): # Skip empty rows
            continue
        # Truncate or pad row to match header count
        padded_row = r[:num_cols] + [""] * (num_cols - len(r))
        cleaned_data_rows.append(padded_row)
        
    if not cleaned_data_rows:
        raise ValueError("No data rows found in the uploaded file.")
        
    # Infer Column Types
    columns_metadata = []
    for col_idx in range(num_cols):
        col_samples = [cleaned_data_rows[row_idx][col_idx] for row_idx in range(min(50, len(cleaned_data_rows)))]
        db_type = infer_postgres_type(col_samples)
        columns_metadata.append((original_headers[col_idx], sanitized_headers[col_idx], db_type))
        
    # 4. Create Isolated PostgreSQL Table
    dataset_id = str(uuid.uuid4())
    safe_uuid = dataset_id.replace('-', '_')
    table_name = f"dataset_{safe_uuid}"
    
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            # Build DDL SQL
            ddl_parts = []
            for _, sanitized, db_type in columns_metadata:
                ddl_parts.append(f'"{sanitized}" {db_type}')
                
            create_sql = f'CREATE TABLE {table_name} ({", ".join(ddl_parts)});'
            cur.execute(create_sql)
            
            # 5. Bulk Insert Rows
            # Build insert query template
            insert_cols = ", ".join(f'"{h}"' for h in sanitized_headers)
            placeholders = ", ".join(["%s"] * num_cols)
            insert_sql = f"INSERT INTO {table_name} ({insert_cols}) VALUES ({placeholders});"
            
            # Map strings to correct inferred types
            typed_rows = []
            for row in cleaned_data_rows:
                typed_row = []
                for col_idx, val in enumerate(row):
                    _, _, db_type = columns_metadata[col_idx]
                    cleaned_val = val.strip()
                    if cleaned_val == "":
                        typed_row.append(None)
                    elif db_type == "INTEGER":
                        try:
                            typed_row.append(int(cleaned_val.replace(',', '')))
                        except ValueError:
                            typed_row.append(None)
                    elif db_type == "DECIMAL(18, 4)":
                        try:
                            typed_row.append(float(cleaned_val.replace(',', '')))
                        except ValueError:
                            typed_row.append(None)
                    else:
                        typed_row.append(val)
                typed_rows.append(typed_row)
                
            cur.executemany(insert_sql, typed_rows)
            
            # Record metadata in datasets_metadata table
            cur.execute("""
                INSERT INTO datasets_metadata (dataset_id, session_id, original_filename, table_name, columns_metadata)
                VALUES (%s, %s, %s, %s, %s);
            """, (dataset_id, session_id, filename, table_name, json.dumps(columns_metadata)))
            
            conn.commit()
            print(f"[DatasetService] Successfully processed and seeded table '{table_name}' for dataset {dataset_id}")
            
            # 6. Index Schema in RAG
            try:
                from app.rag.retriever import index_dataset_schema
                index_dataset_schema(
                    dataset_id=dataset_id,
                    table_name=table_name,
                    filename=filename,
                    columns_metadata=columns_metadata
                )
            except Exception as embed_err:
                print(f"[DatasetService Warning] Initial schema embedding generation deferred: {embed_err}")
                
            return dataset_id
            
    except Exception as e:
        conn.rollback()
        print(f"[DatasetService Error] Failed to seed dataset: {e}")
        # Clean up table if it was created
        try:
            with conn.cursor() as cur:
                cur.execute(f"DROP TABLE IF EXISTS {table_name};")
                conn.commit()
        except Exception:
            pass
        raise e
    finally:
        conn.close()

def delete_user_dataset(dataset_id: str, session_id: str) -> bool:
    """Deletes a custom user dataset and drops its table from PostgreSQL."""
    ensure_metadata_table()
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            # Verify ownership
            cur.execute("""
                SELECT table_name FROM datasets_metadata 
                WHERE dataset_id = %s AND session_id = %s;
            """, (dataset_id, session_id))
            row = cur.fetchone()
            if not row:
                return False
                
            table_name = row[0]
            
            # Drop the table
            cur.execute(f"DROP TABLE IF EXISTS {table_name};")
            # Delete metadata
            cur.execute("DELETE FROM datasets_metadata WHERE dataset_id = %s;", (dataset_id,))
            conn.commit()
            print(f"[DatasetService] Successfully deleted dataset {dataset_id} and table '{table_name}'")
            
            # Clean up RAG schema index and query cache
            try:
                from app.rag.retriever import delete_dataset_schema
                from app.core.cache import query_cache
                delete_dataset_schema(dataset_id)
                query_cache.invalidate_dataset(dataset_id)
            except Exception as e:
                print(f"[DatasetService Warning] Failed to delete schema/query cache: {e}")
                
            return True
    finally:
        conn.close()
